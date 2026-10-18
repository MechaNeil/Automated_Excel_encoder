from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from app.schemas.application import ApplicationData, Person
from app.services.excel_writer import inspect_workbook, write_application


def test_writer_preserves_template_and_writes_adjacent_label(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "UNIT"
    workbook.active["B1"] = "existing"
    workbook.active["A2"] = "FIRST NAME"
    workbook.save(template)
    output = write_application(template, tmp_path / "generated", ApplicationData(unit="Model X", applicant=Person(first_name="Jane")))
    assert inspect_workbook(template)["sheets"] == ["Sheet"]
    assert load_workbook(output).active["B1"].value == "existing"


def test_writer_preserves_template_floating_objects(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "UNIT"
    workbook.save(template)
    with ZipFile(template, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["xl/drawings/drawing1.xml"] = b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" />'
    entries["xl/drawings/_rels/drawing1.xml.rels"] = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships" />'
    entries["xl/media/image1.png"] = b"template-image"
    entries["xl/worksheets/_rels/sheet1.xml.rels"] = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml" /></Relationships>'
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        b"</worksheet>", b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1" /></worksheet>'
    )
    with ZipFile(template, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)

    output = write_application(template, tmp_path / "generated", ApplicationData(unit="Model X"))

    with ZipFile(output) as archive:
        names = set(archive.namelist())
        assert "xl/drawings/drawing1.xml" in names
        assert "xl/media/image1.png" in names
        assert "xl/worksheets/_rels/sheet1.xml.rels" in names
        assert b"drawing1.xml" in archive.read("xl/worksheets/_rels/sheet1.xml.rels")
        assert b"r:id=\"rId1\"" in archive.read("xl/worksheets/sheet1.xml")


def test_writer_writes_to_top_left_cell_of_merged_destination(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "UNIT"
    sheet.merge_cells("B1:C1")
    workbook.save(template)

    output = write_application(template, tmp_path / "generated", ApplicationData(unit="Model X"))

    assert load_workbook(output).active["B1"].value == "Model X"


def test_writer_skips_occupied_adjacent_label(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "UNIT"
    sheet["B1"] = "MAKE/MODEL"
    sheet.merge_cells("B2:C2")
    workbook.save(template)

    output = write_application(template, tmp_path / "generated", ApplicationData(unit="Model X"))

    assert load_workbook(output).active["B2"].value == "Model X"


def test_writer_filters_words_and_separates_borrower_sections(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "PRINCIPAL BORROWER"
    sheet["B1"] = "NAME OF EMPLOYER/BUSINESS"
    sheet["A5"] = "CO-MAKER"
    sheet["B5"] = "EMPLOYER/BUSINESS NAME"
    workbook.save(template)

    output = write_application(
        template,
        tmp_path / "generated",
        ApplicationData(
            applicant=Person(employment_name="Primary Employer"),
            spouse=Person(employment_name="Co-maker Employer"),
        ),
    )

    sheet = load_workbook(output).active
    assert sheet["C1"].value == "Primary Employer"
    assert sheet["C5"].value == "Co-maker Employer"


def test_writer_keeps_age_employer_and_income_fields_aligned(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A7"] = "PRINCIPAL BORROWER"
    sheet["K7"] = "AGE"
    sheet["A21"] = "CO-MAKER"
    sheet["K21"] = "AGE"
    sheet["B14"] = "NAME OF EMPLOYER/BUSINESS:"
    sheet["B28"] = "NAME OF EMPLOYER/BUSINESS:"
    sheet["B35"] = "BORROWER'S MONTHLY SALARY:"
    sheet["B36"] = "SPOUSE'S MONTHLY SALARY:"
    sheet.merge_cells("B14:H14")
    sheet.merge_cells("B28:H28")
    sheet.merge_cells("B35:E35")
    sheet.merge_cells("F35:H35")
    sheet.merge_cells("B36:E36")
    sheet.merge_cells("F36:H36")
    workbook.save(template)

    output = write_application(
        template,
        tmp_path / "generated",
        ApplicationData(
            applicant=Person(age="40", employment_name="Primary Employer", monthly_income="35 - 40K"),
            spouse=Person(age="33", employment_name="Co-maker Employer", monthly_income="18K"),
        ),
    )

    sheet = load_workbook(output).active
    assert sheet["L7"].value == "40"
    assert sheet["L21"].value == "33"
    assert sheet["B15"].value == "Primary Employer"
    assert sheet["B29"].value == "Co-maker Employer"
    assert sheet["F35"].value == "35 - 40K"
    assert sheet["F36"].value == "18K"
