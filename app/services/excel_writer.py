from datetime import datetime
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.schemas.application import ApplicationData
from app.services.field_map import FIELD_ALIASES


def inspect_workbook(path: Path) -> dict:
    if not path.exists():
        return {"exists": False, "path": str(path)}
    workbook = load_workbook(path, data_only=False)
    return {
        "exists": True,
        "sheets": workbook.sheetnames,
        "merged_ranges": {sheet.title: [str(item) for item in sheet.merged_cells.ranges] for sheet in workbook.worksheets},
        "hidden_sheets": [sheet.title for sheet in workbook.worksheets if sheet.sheet_state != "visible"],
        "defined_names": list(workbook.defined_names),
    }


def _label_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def _find_destination(sheet, field: str):
    aliases = {_label_key(alias) for alias in FIELD_ALIASES[field]}
    for row in sheet.iter_rows():
        for cell in row:
            if _label_key(cell.value) in aliases:
                destination = sheet.cell(cell.row, cell.column + 1)
                if isinstance(destination, MergedCell):
                    for merged_range in sheet.merged_cells.ranges:
                        if destination.coordinate in merged_range:
                            destination = sheet.cell(merged_range.min_row, merged_range.min_col)
                            break
                if destination.value:
                    for row_number in range(destination.row + 1, sheet.max_row + 1):
                        candidate = sheet.cell(row_number, destination.column)
                        if not isinstance(candidate, MergedCell) and not candidate.value:
                            return candidate
                return destination
    return None


def write_application(template_path: Path, output_dir: Path, data: ApplicationData) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"Excel template not found: {template_path}")
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template_path)
    values = {"unit": data.unit, "color_first": data.color_first, "color_second": data.color_second}
    for prefix, person in (("applicant", data.applicant), ("spouse", data.spouse)):
        values.update({f"{prefix}.{name}": value for name, value in person.model_dump().items() if name in FIELD_ALIASES})
    written = 0
    for key, value in values.items():
        field = key.split(".")[-1]
        if not value or field not in FIELD_ALIASES:
            continue
        for sheet in workbook.worksheets:
            destination = _find_destination(sheet, field)
            if destination and not destination.value:
                destination.value = value
                written += 1
                break
    if not written:
        raise ValueError("No recognized template labels were found; refusing to generate an unmapped workbook.")
    output = output_dir / f"credit_application_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    workbook.save(output)
    return output
